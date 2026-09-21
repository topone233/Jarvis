from __future__ import annotations

import hashlib
import io
import math
import mimetypes
import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from charset_normalizer import from_bytes

from app.errors import ProviderError, ValidationError
from app.provider import OpenAICompatibleProvider
from app.sections import (
    parse_sections,
    section_label,
)
from app.store import Store
from app.tokens import estimate_tokens
from app.utils import safe_filename, segment_for_index

ALLOWED_EXTENSIONS = {
    ".bat",
    ".c",
    ".cc",
    ".cpp",
    ".cs",
    ".css",
    ".csv",
    ".go",
    ".h",
    ".hpp",
    ".html",
    ".ini",
    ".java",
    ".js",
    ".json",
    ".jsx",
    ".kt",
    ".kts",
    ".md",
    ".markdown",
    ".php",
    ".ps1",
    ".py",
    ".rb",
    ".rs",
    ".sh",
    ".sql",
    ".swift",
    ".toml",
    ".ts",
    ".tsx",
    ".txt",
    ".vue",
    ".xml",
    ".yaml",
    ".yml",
}

#: Documents parsed by a dedicated library rather than decoded as text. The
#: extractors below are loaded lazily so the plain-text path does not pay for
#: their imports at startup.
OFFICE_EXTENSIONS = {
    ".docx",
    ".pdf",
    ".pptx",
    ".xlsx",
}

#: The 97-2003 binary formats. Pure Python cannot read them reliably, so they
#: get their own message instead of the generic unsupported one.
LEGACY_OFFICE_EXTENSIONS = {".doc", ".ppt", ".xls"}

#: A backstop, not a UX limit - the same role MAX_IMAGE_BYTES plays for pastes.
MAX_KNOWLEDGE_BYTES = 50 * 1024 * 1024

#: One spreadsheet row beyond this is truncated, so a worksheet dump cannot
#: turn into an unbounded import.
MAX_SHEET_ROWS = 5_000

#: A single cell holding a novel is pathological; keep it from dominating a chunk.
MAX_CELL_CHARS = 500

#: Total extracted text per document. Chunks are all embedded in one request,
#: so a 100MB extraction would ask the embedding endpoint for thousands of
#: vectors at once - this cap keeps that from happening.
MAX_EXTRACTED_CHARS = 1_000_000

_IMG_TAG_RE = re.compile(r"<img\b[^>]*>")
# Data-URI image markdown mammoth's writer emits; the base64 body would
# otherwise be indexed as noise tokens.
_DATA_URI_IMAGE_RE = re.compile(r"!\[[^\]]*\]\(data:[^)]*\)")
# Any image embed a PDF conversion leaves behind. pymupdf4llm does not write
# image files, so its embeds would point at paths that do not exist.
_IMAGE_EMBED_RE = re.compile(r"!\[[^\]]*\]\([^)]*\)")

# Chars that never carry a meaning worth matching on their own. A query term
# made only of these is the cross-word bigram a bigram index inevitably
# queries (能的 out of 功能的) - coincidence noise that bm25's IDF then
# rewards, because a rare collision outranks a real word. Dropping such terms
# at query time costs nothing: the index keeps every bigram for recall.
_FTS_STOPCHARS = frozenset(
    "的了之乎吗呢吧啊呀哦嘛么和与或及对于从向被给让使在是不没别还就都也很太挺更"
    "最但而且并若如因所由每各该此这那个些会可能应该需要哪何什怎们地得到"
)

#: How many query terms one FTS match may spend.
FTS_MAX_TERMS = 12

# Stated explicitly because Windows resolves mimetypes through the registry,
# which shadows the builtin map and often answers text/plain for OOXML types.
_MIME_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pdf": "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@dataclass(frozen=True)
class ImportItem:
    filename: str
    content: bytes
    relative_path: str | None = None


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    detected = from_bytes(content).best()
    if detected is None:
        raise ValidationError("无法识别文件编码，请仅导入文本或源码文件。")
    return str(detected)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return text[:MAX_CELL_CHARS]


def _docx_to_text(content: bytes) -> str:
    import mammoth

    with io.BytesIO(content) as stream:
        result = mammoth.convert_to_markdown(stream)
    # mammoth turns embedded pictures into base64 data URIs, which would be
    # indexed as noise tokens. The words are what a knowledge base is for.
    return _IMG_TAG_RE.sub("", _DATA_URI_IMAGE_RE.sub("", result.value))


def _xlsx_to_text(content: bytes) -> str:
    from openpyxl import load_workbook

    sections: list[str] = []
    # The stream has to stay open for the whole read: a read_only workbook
    # streams its rows lazily, long after load_workbook has returned.
    with io.BytesIO(content) as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows: list[str] = []
                for row in sheet.iter_rows(values_only=True):
                    cells = [_cell_text(cell) for cell in row]
                    if not any(cells):
                        continue
                    rows.append("| " + " | ".join(cells) + " |")
                    if len(rows) >= MAX_SHEET_ROWS:
                        break
                if rows:
                    sections.append(f"## 表格：{sheet.title}\n" + "\n".join(rows))
        finally:
            workbook.close()
    return "\n\n".join(sections)


def _pptx_to_text(content: bytes) -> str:
    from pptx import Presentation

    with io.BytesIO(content) as stream:
        presentation = Presentation(stream)
    sections: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [_cell_text(cell.text) for cell in row.cells]
                    if any(cells):
                        lines.append("| " + " | ".join(cells) + " |")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(notes)
        if lines:
            sections.append(f"## 幻灯片 {index}\n" + "\n".join(lines))
    return "\n\n".join(sections)


#: How many of a page's first/last nonempty lines count as furniture
#: candidates - the strip a Word export stamps as 页眉/页脚.
_PAGE_EDGE_LINES = 3

#: A line that names nothing but a page number is printer furniture even when
#: every page numbers it differently (第 1 页 共 3 页): the pattern cannot be
#: body text, so it dies at a page edge without needing cross-page agreement.
_PAGE_NUMBER_RE = re.compile(
    r"^(?:第\s*\d+\s*页(?:\s*共\s*\d+\s*页)?|\d+\s*/\s*\d+|page\s+\d+\s*(?:of|/)\s*\d+)$",
    re.IGNORECASE,
)


def _strip_page_furniture(pages: list[str]) -> list[str]:
    """Drop the lines a page layout stamps onto most pages.

    Word exports repeat 页眉/页脚 on every page; landing in one chunk each,
    they polluted embeddings, keyword matches and rerank reads alike. A line
    only counts as furniture when it sits within the first or last few
    nonempty lines of a page AND a majority of pages agree on it - body text
    may repeat, but not at the page edge on most pages. A single-page
    document keeps everything: there is no repetition to observe.
    """
    if len(pages) < 2:
        return pages
    counts: Counter[str] = Counter()
    edges: list[set[str]] = []
    for page in pages:
        lines = [line.strip() for line in page.splitlines() if line.strip()]
        edge = set(lines[:_PAGE_EDGE_LINES] + lines[-_PAGE_EDGE_LINES:])
        edges.append(edge)
        counts.update(edge)
    required = max(2, math.ceil(len(pages) * 0.6))
    furniture = {line for line, count in counts.items() if count >= required}
    furniture |= {line for edge in edges for line in edge if _PAGE_NUMBER_RE.fullmatch(line)}
    if not furniture:
        return pages
    stripped: list[str] = []
    for page in pages:
        kept = [line for line in page.splitlines() if line.strip() not in furniture]
        stripped.append("\n".join(kept).strip("\n"))
    return stripped


def _pdf_to_text(content: bytes) -> str:
    import pymupdf
    import pymupdf4llm

    # Scanned PDFs hold pictures, not text; they come back empty and are
    # skipped with their own message. There is no OCR in this pipeline.
    with pymupdf.open(stream=content, filetype="pdf") as document:
        if not any(page.get_text("text").strip() for page in document):
            return ""
        # to_markdown reads font sizes, so a Word export's larger title lines
        # become real `#` headings the section model can tree into chapters -
        # not one fake `## 第 N 页` per page - and reflows paragraphs. Per-page
        # chunks come back so the 页眉/页脚 a Word export stamps on every page
        # can be stripped before the pages join (pymupdf4llm does not do it).
        pages = _strip_page_furniture(
            [chunk["text"] for chunk in pymupdf4llm.to_markdown(document, page_chunks=True)]
        )
    # write_images never writes files, but image blocks can still come through
    # as markdown embeds pointing nowhere. The words are what a knowledge
    # base is for - same rule the docx path applies to data-URI pictures.
    return _IMAGE_EMBED_RE.sub("", "\n".join(pages)).strip()


_EXTRACTORS = {
    ".docx": _docx_to_text,
    ".pdf": _pdf_to_text,
    ".pptx": _pptx_to_text,
    ".xlsx": _xlsx_to_text,
}


def _extract_text(suffix: str, content: bytes) -> str:
    extractor = _EXTRACTORS.get(suffix)
    if extractor is None:
        return _decode_text(content)
    try:
        text = extractor(content)
    except ValidationError:
        raise
    except Exception as error:
        # A mislabeled file (a renamed text file with a .docx suffix, a
        # truncated download) arrives as arbitrary bytes; the parsing
        # libraries raise half a dozen different exceptions for that, so the
        # catch has to be wide. The message is all the caller sees.
        raise ValidationError("无法解析这个文档，文件可能已损坏或不是它声称的格式。") from error
    return text[:MAX_EXTRACTED_CHARS]


def _split_long_text(value: str, target_size: int, overlap_size: int) -> list[str]:
    chunks: list[str] = []
    start = 0
    while start < len(value):
        end = min(len(value), start + target_size)
        if end < len(value):
            split_at = value.rfind("\n", start + target_size // 2, end)
            if split_at > start:
                end = split_at
        part = value[start:end].strip()
        if part:
            chunks.append(part)
        if end == len(value):
            break
        start = max(end - overlap_size, start + 1)
    return chunks


def chunk_text(
    content: str,
    target_size: int = 1_800,
    overlap_size: int = 220,
    fallback_title: str = "全文",
) -> list[dict[str, str]]:
    """One chunk per section, never crossing a section boundary.

    A test-case document's sections are exactly what a query matches against,
    so the section - not a blank-line window - is the retrieval unit: it lets
    a hit be cited down to its own heading, and it stops a chunk from blending
    half of one case with half of another. A section's chunk covers its *own*
    text only - everything before its first child heading, or the whole span
    for a leaf - because a parent's full span would re-embed every descendant,
    and the big duplicated chunks would crowd the real hits out of the top-k.
    A heading that only groups children (no text of its own) gets no chunk;
    its words survive in the children's breadcrumbs. Oversized sections still
    split internally (same long-text rule as before); every piece keeps the
    section's id and opens with its breadcrumb, which is what carries the
    heading's words into both the FTS index and the embedding.
    """
    sections = parse_sections(content, fallback_title)
    by_id = {section.id: section for section in sections}
    first_child: dict[str, int] = {}
    for section in sections:
        if "." in section.id:
            parent = section.id.rsplit(".", 1)[0]
            start = first_child.get(parent)
            first_child[parent] = section.start if start is None else min(start, section.start)
    chunks: list[dict[str, str]] = []
    for section in sections:
        own_end = first_child.get(section.id, section.end)
        # The body starts past the section's own heading line - but only when
        # there is one: a headingless document's root section begins at the
        # first word, and skipping its first line would silently drop it. A
        # heading that only groups children then has an empty body and gets
        # no chunk (its words live in the children's breadcrumbs), and a
        # leaf's text is its content without the heading the breadcrumb
        # already restates.
        newline = content.find("\n", section.start, own_end)
        first_line_end = own_end if newline == -1 else newline
        starts_with_heading = content[section.start : first_line_end].lstrip().startswith("#")
        body_start = section.start if not starts_with_heading else first_line_end + 1
        text = content[body_start:own_end].strip()
        if not text:
            continue
        header = f"【{section_label(section, by_id)}】\n"
        if len(text) > target_size:
            parts = _split_long_text(text, target_size, overlap_size)
        else:
            parts = [text]
        for part in parts:
            chunks.append({"content": f"{header}{part}", "section_id": section.id})
    return chunks


class KnowledgeService:
    #: How many mixed hits a rerank pass sees. Its whole value is picking
    #: better from a wider net, so when one is configured the hybrid stage
    #: casts further than the caller asked and the reranker does the cutting.
    RERANK_CANDIDATES = 20

    def __init__(self, store: Store, provider: OpenAICompatibleProvider) -> None:
        self.store = store
        self.provider = provider

    async def import_items(
        self,
        items: list[ImportItem],
        *,
        project_id: str | None,
    ) -> list[dict[str, Any]]:
        if not items:
            raise ValidationError("请至少选择一个文件。")
        embedding_spec = self.store.get_retrieval_spec("embedding")
        results: list[dict[str, Any]] = []
        for item in items:
            suffix = Path(item.filename).suffix.lower()
            if suffix in LEGACY_OFFICE_EXTENSIONS:
                results.append(
                    {
                        "filename": item.filename,
                        "status": "skipped",
                        "reason": "旧版 Office 格式暂不支持，请在 Office 里另存为 "
                        ".docx/.xlsx/.pptx 后再导入。",
                    }
                )
                continue
            if suffix not in ALLOWED_EXTENSIONS and suffix not in OFFICE_EXTENSIONS:
                results.append(
                    {
                        "filename": item.filename,
                        "status": "skipped",
                        "reason": "当前仅支持文本、Markdown、源码和 docx/xlsx/pptx/pdf 文档。",
                    }
                )
                continue
            if not item.content:
                results.append(
                    {"filename": item.filename, "status": "skipped", "reason": "文件为空。"}
                )
                continue
            if len(item.content) > MAX_KNOWLEDGE_BYTES:
                results.append(
                    {
                        "filename": item.filename,
                        "status": "skipped",
                        "reason": "单个文件不能超过 50MB。",
                    }
                )
                continue
            try:
                text = _extract_text(suffix, item.content)
            except ValidationError as error:
                results.append(
                    {"filename": item.filename, "status": "skipped", "reason": str(error)}
                )
                continue
            if suffix in OFFICE_EXTENSIONS and not text.strip():
                reason = (
                    "这份 PDF 没有可提取的文字（可能是扫描件）。"
                    if suffix == ".pdf"
                    else "未在文档中找到可提取的文字。"
                )
                results.append({"filename": item.filename, "status": "skipped", "reason": reason})
                continue
            content_hash = hashlib.sha256(item.content).hexdigest()
            document_id_hint = hashlib.sha1(
                f"{item.filename}:{content_hash}".encode(), usedforsecurity=False
            ).hexdigest()[:12]
            stored_filename = f"{document_id_hint}_{safe_filename(item.filename)}"
            stored_path = self.store.database.objects_directory / stored_filename
            stored_path.write_bytes(item.content)
            document = self.store.create_knowledge_document(
                project_id=project_id,
                title=Path(item.filename).stem or item.filename,
                original_filename=item.filename,
                relative_path=item.relative_path,
                mime_type=_MIME_TYPES.get(suffix)
                or mimetypes.guess_type(item.filename)[0]
                or "text/plain",
                content_hash=content_hash,
                stored_path=str(stored_path),
                content=text,
            )
            parts = chunk_text(text, fallback_title=Path(item.filename).stem or item.filename)
            payloads: list[dict[str, Any]] = [
                {
                    "position": index,
                    "content": part["content"],
                    "section_id": part["section_id"],
                    "token_estimate": estimate_tokens(part["content"]),
                }
                for index, part in enumerate(parts)
            ]
            status = "ready"
            embedding_model: str | None = None
            if embedding_spec:
                try:
                    vectors = await self.provider.embed(
                        embedding_spec, [part["content"] for part in payloads]
                    )
                    if len(vectors) != len(payloads):
                        raise ProviderError("嵌入服务返回的向量数量不匹配。")
                    for payload, vector in zip(payloads, vectors, strict=True):
                        payload["embedding"] = vector
                    embedding_model = embedding_spec["model"]
                except ProviderError as error:
                    status = "ready_without_embeddings"
                    results.append(
                        {
                            "filename": item.filename,
                            "status": "warning",
                            "reason": f"已完成关键词索引，语义索引暂不可用：{error}",
                        }
                    )
            self.store.create_knowledge_chunks(
                document["id"], project_id, payloads, embedding_model
            )
            completed = self.store.finish_knowledge_document(document["id"], status, len(payloads))
            results.append({"filename": item.filename, "status": status, "document": completed})
        return results

    async def search(
        self,
        query: str,
        *,
        project_id: str | None,
        limit: int = 8,
        query_vector: list[float] | None = None,
    ) -> list[dict[str, Any]]:
        """Hybrid keyword + semantic hits for one query, optionally reranked.

        `query_vector` lets a caller that already embedded this exact query
        hand the vector over - the context build ranks memories with it
        first - instead of paying for a second identical embedding call. With
        a rerank model configured, the mixed ranking only nominates
        RERANK_CANDIDATES; the reranker picks the final `limit` and its
        relevance score becomes the hit's score.
        """
        embedding_spec = self.store.get_retrieval_spec("embedding")
        rerank_spec = self.store.get_retrieval_spec("rerank")
        candidate_limit = self.RERANK_CANDIDATES if rerank_spec else limit
        fts_terms = self._fts_terms(query)
        fts_results = (
            self.store.search_knowledge_fts(
                " OR ".join(f'"{term}"' for term in fts_terms), project_id
            )
            if fts_terms
            else []
        )
        scored: dict[str, tuple[float, dict[str, Any]]] = {}
        fts_ids: set[str] = set()
        # FTS rank alone says nothing: bm25 hands its top slot to a chunk that
        # matches one rare coincidence bigram (能的 out of 性能的) even when the
        # query has nothing to do with it, and the old position-only weight
        # turned that junk into a 0.45 citation. A hit's score is weighted by
        # how much of the query it actually matched, and when the query offers
        # more than one term a hit must match at least two - a real hit nearly
        # always does, and the single-collision junk drops out entirely.
        total_terms = len(fts_terms)
        for index, item in enumerate(fts_results):
            matched = sum(1 for term in fts_terms if term.casefold() in item["content"].casefold())
            if total_terms >= 2 and matched < 2:
                continue
            fts_ids.add(item["id"])
            scored[item["id"]] = (0.45 * matched / total_terms / (index + 1), item)
        if embedding_spec:
            try:
                if query_vector is None:
                    query_vector = (await self.provider.embed(embedding_spec, [query]))[0]
                vec_hits = self.store.search_knowledge_vec(
                    project_id,
                    embedding_spec["model"],
                    query_vector,
                    candidate_limit,
                )
                for item in vec_hits:
                    # Cosine distance to similarity, on the same 0.75 weight
                    # the Python scan used, so ranking contracts stay put.
                    # None is a zero vector's NaN - sqlite-vec hands NaN back
                    # as NULL - and means "no similarity", the same verdict
                    # _cosine gives a zero vector, not a crash.
                    distance = item["distance"]
                    score = 0.0 if distance is None else max(1.0 - distance, 0.0)
                    score *= 0.75
                    existing = scored.get(item["id"])
                    scored[item["id"]] = (score + (existing[0] if existing else 0), item)
            except (ProviderError, IndexError):
                pass
        candidates = sorted(scored.values(), key=lambda item: item[0], reverse=True)[
            :candidate_limit
        ]
        reranked = False
        if rerank_spec and len(candidates) > 1:
            try:
                ranking = await self.provider.rerank(
                    rerank_spec,
                    query,
                    [item["content"] for _, item in candidates],
                    limit,
                )
                candidates = [(relevance, candidates[index][1]) for index, relevance in ranking]
                reranked = True
            except ProviderError:
                # A dead reranker must not take retrieval down with it: the
                # mixed order stands, only the quality of the cut suffers.
                pass
        if not reranked:
            candidates = candidates[:limit]
        # A hit's section label comes from the section model of the document
        # it lives in. Documents are parsed once per call, however many of
        # their chunks landed in the top-k.
        section_models: dict[str, tuple[str, dict[str, Any]]] = {}
        results: list[dict[str, Any]] = []
        for score, item in candidates:
            section_id = item.get("section_id")
            section_title = None
            if section_id:
                model = section_models.get(item["document_id"])
                if model is None:
                    document = self.store.get_knowledge_document(item["document_id"])
                    sections = parse_sections(document["content"], document["title"])
                    model = (document["content"], {s.id: s for s in sections})
                    section_models[item["document_id"]] = model
                label = section_label(model[1][section_id], model[1])
                section_title = label
            results.append(
                {
                    "chunk_id": item["id"],
                    "document_id": item["document_id"],
                    "title": item.get("document_title", "未命名文档"),
                    "section_id": section_id,
                    "section_title": section_title,
                    "content": item["content"],
                    "score": round(score, 4),
                    "source": "reranked"
                    if reranked
                    else ("hybrid" if item["id"] in fts_ids else "semantic"),
                }
            )
        return results

    @staticmethod
    def _fts_terms(query: str) -> list[str]:
        """The query's FTS terms: bigram-segmented words, deduplicated
        case-insensitively, capped - and minus the terms made only of
        stopword chars, which can only ever match by coincidence."""
        terms: list[str] = []
        seen: set[str] = set()
        for term in re.findall(r"[\w\u4e00-\u9fff]+", segment_for_index(query)):
            folded = term.casefold()
            if folded in seen or all(char in _FTS_STOPCHARS for char in term):
                continue
            seen.add(folded)
            terms.append(term)
            if len(terms) == FTS_MAX_TERMS:
                break
        return terms

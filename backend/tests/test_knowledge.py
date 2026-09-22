from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from typing import Any

import pytest

from app import knowledge as knowledge_module
from app.knowledge import ImportItem
from app.runtime import CoreServices
from app.sections import parse_sections
from app.utils import segment_for_index


@pytest.mark.asyncio
async def test_import_and_hybrid_search(
    core: CoreServices, use_embedding: Callable[[], None]
) -> None:
    use_embedding()
    project = core.store.create_project("Jarvis", False)
    imported = await core.knowledge.import_items(
        [
            ImportItem(
                filename="architecture.md",
                content="Jarvis 使用 SQLite 保存本地知识，并支持上下文压缩。".encode(),
                relative_path="notes/architecture.md",
            )
        ],
        project_id=project["id"],
    )

    assert imported[-1]["status"] == "ready"
    document = imported[-1]["document"]
    assert document["chunk_count"] == 1

    results = await core.knowledge.search("SQLite 上下文", project_id=project["id"])

    assert results
    assert results[0]["document_id"] == document["id"]
    assert "SQLite" in results[0]["content"]


@pytest.mark.asyncio
async def test_chinese_query_matches_a_sub_phrase(core: CoreServices) -> None:
    """Keyword search must reach inside a CJK run, not only match whole runs."""
    project = core.store.create_project("中文检索", False)
    await core.knowledge.import_items(
        [
            ImportItem(
                filename="notes.md",
                content="Jarvis 使用 SQLite 保存本地知识，并支持上下文压缩。".encode(),
            )
        ],
        project_id=project["id"],
    )

    for query in ("上下文", "压缩", "上下文压缩"):
        results = await core.knowledge.search(query, project_id=project["id"])
        assert results, f"未命中：{query}"

    assert not await core.knowledge.search("记忆", project_id=project["id"])


def test_segment_for_index_expands_cjk_runs_into_bigrams() -> None:
    assert segment_for_index("上下文") == "上下 下文"
    assert segment_for_index("记忆") == "记忆"
    assert segment_for_index("SQLite 上下文压缩") == "SQLite 上下 下文 文压 压缩"
    assert segment_for_index("plain english text") == "plain english text"


@pytest.mark.asyncio
async def test_import_never_writes_the_source_file(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [ImportItem(filename="note.txt", content=b"original content")], project_id=None
    )
    document = imported[-1]["document"]

    assert document["stored_path"].startswith(str(core.database.objects_directory))
    assert core.database.objects_directory.joinpath(
        document["stored_path"].split("\\")[-1]
    ).exists()


@pytest.mark.asyncio
async def test_import_stores_content_and_section_shaped_chunks(core: CoreServices) -> None:
    content = "# 指南\n## 登录\n账号密码登录的用例。\n## 上传\n附件上传的用例。\n"
    imported = await core.knowledge.import_items(
        [ImportItem(filename="guide.md", content=content.encode())], project_id=None
    )
    document = imported[-1]["document"]

    stored = core.store.get_knowledge_document(document["id"])
    assert stored["content"] == content

    chunks = core.database.fetchall(
        "SELECT section_id, content FROM knowledge_chunks"
        " WHERE document_id = ? AND deleted_at IS NULL ORDER BY position",
        (document["id"],),
    )
    assert [row["section_id"] for row in chunks] == ["1.1", "1.2"]
    assert chunks[0]["content"].startswith("【1.1 登录】")
    assert "账号密码" in chunks[0]["content"]


@pytest.mark.asyncio
async def test_search_reports_the_section_of_a_hit(core: CoreServices) -> None:
    content = "# doc\n## 用户管理\n登录的测试用例。\n## 订单管理\n下单的测试用例。\n"
    imported = await core.knowledge.import_items(
        [ImportItem(filename="cases.md", content=content.encode())], project_id=None
    )
    document = imported[-1]["document"]

    results = await core.knowledge.search("登录 测试用例", project_id=None)
    assert results
    assert results[0]["document_id"] == document["id"]
    assert results[0]["section_id"] == "1.1"
    assert "用户管理" in results[0]["section_title"]


# --- Document fixtures -------------------------------------------------------
#
# Hand-built rather than checked-in binaries, so every test states the file
# shape it depends on and none of them age out of sync with a blob.


def _zip_of(files: dict[str, str | bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, payload in files.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def _docx(paragraphs: list[str], *, with_image: bool = False) -> bytes:
    """A minimal .docx: the parts mammoth needs and nothing else.

    With `with_image`, a paragraph also carries a drawing whose blip
    references rId5, so a test can prove the extracted text keeps only words.
    """
    drawing = (
        "<w:p><w:r><w:drawing>"
        '<wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        '<pic:blipFill><a:blip r:embed="rId5"/></pic:blipFill>'
        "</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing></w:r></w:p>"
        if with_image
        else ""
    )
    body = drawing + "".join(f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" for text in paragraphs)
    return _zip_of(
        {
            "[Content_Types].xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                '<Default Extension="rels" ContentType='
                '"application/vnd.openxmlformats-package.relationships+xml"/>'
                '<Default Extension="xml" ContentType="application/xml"/>'
                '<Default Extension="png" ContentType="image/png"/>'
                '<Override PartName="/word/document.xml" ContentType='
                '"application/vnd.openxmlformats-officedocument.wordprocessingml'
                '.document.main+xml"/>'
                "</Types>"
            ),
            "_rels/.rels": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type='
                '"http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
                'officeDocument" Target="word/document.xml"/></Relationships>'
            ),
            "word/_rels/document.xml.rels": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId5" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                'Target="media/image1.png"/></Relationships>'
            ),
            "word/media/image1.png": b"\x89PNG-fixture-bytes",
            "word/document.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
                ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                f"<w:body>{body}</w:body></w:document>"
            ),
        }
    )


def _xlsx(sheets: dict[str, list[list[Any]]]) -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    # A fresh workbook always has one sheet; the stubs hedge, so narrow it.
    first = workbook.active
    if first is not None:
        workbook.remove(first)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "发布会时间表"
    second = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    second.text_frame.text = "Quarterly roadmap review"
    slide.notes_slide.notes_text_frame.text = "备注：先讲预算"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _pdf(pages: list[str], *, font: str = "Helvetica") -> bytes:
    """A minimal PDF with one text line per page, xref offsets computed for real."""
    objects: list[bytes | None] = []

    def add(body: bytes | None) -> int:
        objects.append(body)
        return len(objects)

    catalog_id = add(None)
    pages_id = add(None)
    if font == "Helvetica":
        font_object = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    else:
        # A non-embedded CID font: read through the reader's bundled CMaps,
        # so a Chinese fixture needs no font file. MuPDF refuses a CID font
        # without a FontDescriptor, so a minimal one rides along.
        descriptor_id = add(
            b"<< /Type /FontDescriptor /FontName /STSong-Light /Flags 4 "
            b"/FontBBox [0 0 1000 1000] /ItalicAngle 0 /Ascent 800 /Descent -200 "
            b"/CapHeight 700 /StemV 80 >>"
        )
        font_object = (
            b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /UniGB-UCS2-H "
            b"/DescendantFonts [<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light "
            b"/FontDescriptor " + str(descriptor_id).encode() + b" 0 R "
            b"/CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 4 >> >>] >>"
        )
    font_id = add(font_object)
    kids: list[int] = []
    for text in pages:
        if font == "Helvetica":
            encoded = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        else:
            # UniGB-UCS2-H takes UTF-16BE code units as hex strings.
            hex_text = text.encode("utf-16-be").hex().upper()
            encoded = f"BT /F1 12 Tf 72 720 Td <{hex_text}> Tj ET".encode()
        stream_id = add(
            b"<< /Length "
            + str(len(encoded)).encode()
            + b" >>\nstream\n"
            + encoded
            + b"\nendstream"
        )
        page = (
            f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {stream_id} 0 R >>"
        ).encode()
        kids.append(add(page))
    objects[catalog_id - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode()
    objects[pages_id - 1] = (
        f"<< /Type /Pages /Kids [{' '.join(f'{kid} 0 R' for kid in kids)}] /Count {len(kids)} >>"
    ).encode()
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for object_id, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{object_id} 0 obj\n".encode() + (body or b"<< >>") + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF"
    ).encode()
    return bytes(out)


# --- Office and PDF imports --------------------------------------------------


@pytest.mark.asyncio
async def test_docx_import_extracts_paragraphs(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [ImportItem(filename="设计.docx", content=_docx(["Jarvis 的知识库设计", "Local first"]))],
        project_id=None,
    )
    assert imported[-1]["status"] == "ready"
    document = imported[-1]["document"]
    assert document["chunk_count"] == 1

    results = await core.knowledge.search("知识库", project_id=None)
    assert results
    assert results[0]["document_id"] == document["id"]


@pytest.mark.asyncio
async def test_docx_import_keeps_words_but_not_picture_bytes(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [ImportItem(filename="带图.docx", content=_docx(["知识库正文"], with_image=True))],
        project_id=None,
    )
    assert imported[-1]["status"] == "ready"

    results = await core.knowledge.search("知识库正文", project_id=None)
    assert results
    assert "data:image" not in results[0]["content"]


@pytest.mark.asyncio
async def test_xlsx_import_extracts_sheet_rows(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [
            ImportItem(
                filename="预算.xlsx",
                content=_xlsx({"二月": [["项目", "金额"], ["服务器", 3200], ["会议", 500]]}),
            )
        ],
        project_id=None,
    )
    assert imported[-1]["status"] == "ready"

    results = await core.knowledge.search("服务器", project_id=None)
    assert results
    assert "3200" in results[0]["content"]


@pytest.mark.asyncio
async def test_xlsx_import_truncates_huge_sheets(core: CoreServices) -> None:
    rows = [["项目", "金额"]] + [[f"值{index}", index] for index in range(5_100)]
    imported = await core.knowledge.import_items(
        [ImportItem(filename="大表.xlsx", content=_xlsx({"流水": rows}))], project_id=None
    )
    assert imported[-1]["status"] == "ready"

    hits = await core.knowledge.search("值4998", project_id=None)
    assert hits
    assert not await core.knowledge.search("值4999", project_id=None)


@pytest.mark.asyncio
async def test_pptx_import_extracts_slides_and_notes(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [ImportItem(filename="评审.pptx", content=_pptx())], project_id=None
    )
    assert imported[-1]["status"] == "ready"

    results = await core.knowledge.search("roadmap", project_id=None)
    assert results
    assert "备注：先讲预算" in results[0]["content"]


@pytest.mark.asyncio
async def test_pdf_import_extracts_pages(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [
            ImportItem(
                filename="notes.pdf", content=_pdf(["Jarvis knowledge base", "retrieval page two"])
            )
        ],
        project_id=None,
    )
    assert imported[-1]["status"] == "ready"

    results = await core.knowledge.search("retrieval", project_id=None)
    assert results
    # Pages flow into one body with no `## 第 N 页` headings - the fake page
    # structure was what broke the section model on every PDF import.
    content = results[0]["content"]
    assert "Jarvis knowledge base" in content
    assert "retrieval page two" in content
    assert "第" not in content


@pytest.mark.asyncio
async def test_pdf_import_reads_chinese_through_a_cid_font(core: CoreServices) -> None:
    """A non-embedded CID font must carry Chinese text without a font file."""
    imported = await core.knowledge.import_items(
        [ImportItem(filename="中文.pdf", content=_pdf(["知识库存放本地文档"], font="STSong"))],
        project_id=None,
    )
    assert imported[-1]["status"] == "ready"

    results = await core.knowledge.search("知识库", project_id=None)
    assert results
    assert "知识库存放本地文档" in results[0]["content"]


@pytest.mark.asyncio
async def test_pdf_headings_come_from_font_sizes_not_page_breaks(core: CoreServices) -> None:
    """A Word-style export's larger title lines become real headings, so the
    section model trees the document by its own structure."""
    import pymupdf

    buffer = io.BytesIO()
    with pymupdf.open() as source:
        page = source.new_page()
        page.insert_text((72, 96), "Login Module Design", fontsize=20)
        page.insert_text((72, 140), "Login Endpoint", fontsize=14)
        page.insert_text((72, 160), "The login endpoint validates a captcha.", fontsize=11)
        page.insert_text((72, 190), "Token Endpoint", fontsize=14)
        page.insert_text((72, 210), "The token endpoint signs a session.", fontsize=11)
        source.save(buffer)
    imported = await core.knowledge.import_items(
        [ImportItem(filename="login.pdf", content=buffer.getvalue())], project_id=None
    )
    assert imported[-1]["status"] == "ready"
    document = imported[-1]["document"]

    # Two heading tiers become a chapter with two sections - the tree the
    # tool navigates, instead of one `## 第 N 页` per page.
    sections = parse_sections(document["content"], "全文")
    assert [section.title for section in sections] == [
        "Login Module Design",
        "Login Endpoint",
        "Token Endpoint",
    ]
    assert [section.id for section in sections] == ["1", "1.1", "1.2"]


@pytest.mark.asyncio
async def test_pdf_strips_repeated_page_furniture(core: CoreServices) -> None:
    """A Word export stamps 页眉 and a per-page 页码 on every page; into the
    knowledge base they would ride every chunk. Repetition across pages and
    the page-number pattern both die; the real content survives."""
    import pymupdf

    buffer = io.BytesIO()
    with pymupdf.open() as source:
        for index in range(1, 4):
            page = source.new_page()
            page.insert_text((72, 40), "Jarvis 项目内部设计文档", fontsize=9, fontname="china-s")
            page.insert_text((72, 780), f"第 {index} 页 共 3 页", fontsize=9, fontname="china-s")
            page.insert_text(
                (72, 120 + 20 * index),
                f"登录接口的验收标准第{index}条：验证码校验。",
                fontsize=11,
                fontname="china-s",
            )
        source.save(buffer)
    imported = await core.knowledge.import_items(
        [ImportItem(filename="导出.pdf", content=buffer.getvalue())], project_id=None
    )
    assert imported[-1]["status"] == "ready"
    content = imported[-1]["document"]["content"]

    assert "项目内部设计文档" not in content
    assert "共 3 页" not in content
    assert "验收标准第1条" in content
    assert "验收标准第3条" in content


@pytest.mark.asyncio
async def test_textless_pdf_is_reported_as_scanned(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [ImportItem(filename="扫描.pdf", content=_pdf([]))], project_id=None
    )
    assert imported[-1]["status"] == "skipped"
    assert "扫描件" in imported[-1]["reason"]


@pytest.mark.asyncio
async def test_legacy_office_formats_ask_for_a_resave(core: CoreServices) -> None:
    imported = await core.knowledge.import_items(
        [ImportItem(filename="旧文档.doc", content=b"whatever")], project_id=None
    )
    assert imported[-1]["status"] == "skipped"
    assert "另存为" in imported[-1]["reason"]


@pytest.mark.asyncio
async def test_mislabeled_documents_are_skipped_not_fatal(core: CoreServices) -> None:
    results = await core.knowledge.import_items(
        [
            ImportItem(filename="假的.docx", content=b"this is not a zip file"),
            ImportItem(filename="截断.pdf", content=b"%PDF-1.4\n"),
        ],
        project_id=None,
    )
    assert all(result["status"] == "skipped" for result in results)
    assert all(result["reason"] for result in results)


@pytest.mark.asyncio
async def test_import_rejects_files_over_the_backstop(
    core: CoreServices, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(knowledge_module, "MAX_KNOWLEDGE_BYTES", 10)
    imported = await core.knowledge.import_items(
        [ImportItem(filename="太大.txt", content=b"x" * 11)], project_id=None
    )
    assert imported[-1]["status"] == "skipped"
    assert "50MB" in imported[-1]["reason"]


# --- sqlite-vec KNN index ----------------------------------------------------


@pytest.mark.asyncio
async def test_semantic_search_hits_the_nearest_chunk(
    core: CoreServices, use_embedding: Callable[[], None]
) -> None:
    """A query whose vector exactly equals one chunk's must rank it first."""
    use_embedding()
    short = core.knowledge.import_items(
        [ImportItem(filename="短.txt", content=b" precisely seventeen ")], project_id=None
    )
    long_import = core.knowledge.import_items(
        [
            ImportItem(
                filename="长.txt",
                content=b"a far longer passage with many more words in it",
            )
        ],
        project_id=None,
    )
    short_doc = (await short)[-1]["document"]
    long_doc = (await long_import)[-1]["document"]

    results = await core.knowledge.search(" precisely seventeen ", project_id=None)
    assert results
    assert results[0]["document_id"] == short_doc["id"]
    assert results[0]["score"] > 0.5

    others = await core.knowledge.search(
        "a far longer passage with many more words in it", project_id=None
    )
    assert others
    assert others[0]["document_id"] == long_doc["id"]


@pytest.mark.asyncio
async def test_delete_removes_the_vec_rows(
    core: CoreServices, use_embedding: Callable[[], None]
) -> None:
    use_embedding()
    imported = await core.knowledge.import_items(
        [ImportItem(filename="临时.txt", content="被删除后就不再被向量检索命中".encode())],
        project_id=None,
    )
    document = imported[-1]["document"]
    before = core.store.database.fetchall(
        "SELECT COUNT(*) AS n FROM knowledge_chunks_vec_2 WHERE chunk_id IN"
        " (SELECT id FROM knowledge_chunks WHERE document_id = ?)",
        (document["id"],),
    )
    assert before[0]["n"] > 0

    core.store.delete_knowledge_document(document["id"])
    after = core.store.database.fetchall(
        "SELECT COUNT(*) AS n FROM knowledge_chunks_vec_2 WHERE chunk_id IN"
        " (SELECT id FROM knowledge_chunks WHERE document_id = ?)",
        (document["id"],),
    )
    assert after[0]["n"] == 0


@pytest.mark.asyncio
async def test_restore_rebuilds_fts_and_vec_from_chunks(
    core: CoreServices, use_embedding: Callable[[], None]
) -> None:
    """Deleting drops the derived indexes; restoring must bring both back."""
    use_embedding()
    imported = await core.knowledge.import_items(
        [ImportItem(filename="复活.txt", content="恢复之后要能被重新检索到".encode())],
        project_id=None,
    )
    document = imported[-1]["document"]
    core.store.delete_knowledge_document(document["id"])
    assert not await core.knowledge.search("恢复", project_id=None)

    trash = core.store.list_trash()
    core.store.restore_trash_item(trash[0]["id"])

    results = await core.knowledge.search("恢复", project_id=None)
    assert results
    assert results[0]["document_id"] == document["id"]


@pytest.mark.asyncio
async def test_a_missing_vec_table_means_no_semantic_hits(core: CoreServices) -> None:
    """A width nobody has embedded at yet is empty, not an error."""
    hits = core.store.search_knowledge_vec(None, "mock-embedding", [0.5] * 7, limit=5)
    assert hits == []


@pytest.mark.asyncio
async def test_a_zero_vector_is_no_similarity_not_a_crash(
    core: CoreServices, use_embedding: Callable[[], None], monkeypatch: pytest.MonkeyPatch
) -> None:
    """A chunk that embeds to the zero vector gets NaN from cosine, which
    sqlite-vec hands back as a NULL distance. That is "no similarity" - the
    same verdict a zero vector gets everywhere else - not a failed search."""

    async def zero(spec: dict[str, str], texts: list[str]) -> list[list[float]]:
        del spec, texts
        return [[0.0, 0.0]]

    use_embedding()
    monkeypatch.setattr(core.knowledge.provider, "embed", zero)
    await core.knowledge.import_items(
        [ImportItem(filename="零.txt", content="没有任何特征的内容".encode())],
        project_id=None,
    )

    results = await core.knowledge.search("内容", project_id=None)

    assert results
    # The semantic weight (0.75) contributed nothing; what is left is the
    # keyword side of the mix.
    assert results[0]["score"] <= 0.45


@pytest.mark.asyncio
async def test_a_single_term_collision_no_longer_becomes_a_citation(core: CoreServices) -> None:
    """性能的… carries the 能的 bigram a 功能的… query asks for. bm25 happily
    ranked that coincidence first and the position-only weight turned it
    into a 0.45 citation; a multi-term query now demands two matches, and
    a query made only of stopword chars asks for nothing at all."""
    await core.knowledge.import_items(
        [ImportItem(filename="性能对比.md", content="性能的性能对比数据。".encode())],
        project_id=None,
    )
    assert not await core.knowledge.search("功能的测试", project_id=None)
    assert not await core.knowledge.search("的了吗", project_id=None)


@pytest.mark.asyncio
async def test_keyword_score_weights_how_much_of_the_query_matched(core: CoreServices) -> None:
    """Full coverage keeps the 0.45 ceiling; the half-matched chunk gets its
    coverage fraction (times its bm25 rank decay, rank 2 here); the chunk
    whose only tie was the dropped 能的 collision is not cited at all. Chunk
    content opens with its breadcrumb, so the text assertions look inside."""
    await core.knowledge.import_items(
        [ImportItem(filename="性能对比.md", content="性能的性能对比数据。".encode())],
        project_id=None,
    )
    await core.knowledge.import_items(
        [ImportItem(filename="功能说明.md", content="功能模块覆盖测试。".encode())],
        project_id=None,
    )
    await core.knowledge.import_items(
        [ImportItem(filename="登录用例.md", content="登录功能的测试用例：弱口令。".encode())],
        project_id=None,
    )
    results = await core.knowledge.search("功能的测试", project_id=None)
    assert [result["title"] for result in results] == ["登录用例", "功能说明"]
    assert "登录功能的测试用例" in results[0]["content"]
    assert "功能模块覆盖测试" in results[1]["content"]
    assert results[0]["score"] == 0.45
    assert results[1]["score"] == 0.15


@pytest.mark.asyncio
async def test_migration_backfills_vec_rows_from_stored_embeddings(
    core: CoreServices, use_embedding: Callable[[], None]
) -> None:
    """A database from schema 4 has embeddings but no vec rows; backfill them."""
    use_embedding()
    imported = await core.knowledge.import_items(
        [ImportItem(filename="旧库.txt", content="迁移之后向量行应该被补齐".encode())],
        project_id=None,
    )
    document = imported[-1]["document"]
    with core.database.transaction() as connection:
        connection.execute("DELETE FROM knowledge_chunks_vec_2")
        core.database._backfill_vec_index(connection)
    rows = core.store.database.fetchall(
        "SELECT COUNT(*) AS n FROM knowledge_chunks_vec_2 WHERE chunk_id IN"
        " (SELECT id FROM knowledge_chunks WHERE document_id = ?)",
        (document["id"],),
    )
    assert rows[0]["n"] > 0


def test_strip_page_furniture_counts_folio_lines_with_digits_folded() -> None:
    """A header that carries its own page number never repeats verbatim; it
    is folio-shaped, so folding digits makes 「第 3 页」 and 「第 4 页」 one
    line. A body line whose digits vary is not folio-shaped and survives."""
    pages = [
        f"产品手册 · 第 {index} 页\n\n登录模块第{index}章的接口约定。\n第{index}节描述操作步骤。"
        for index in range(1, 4)
    ]
    stripped = knowledge_module._strip_page_furniture(pages)
    joined = "\n".join(stripped)
    assert "产品手册" not in joined
    assert "登录模块第1章的接口约定。" in joined
    assert "登录模块第3章的接口约定。" in joined
    assert "第3节描述操作步骤。" in joined


def test_strip_page_furniture_kills_decorated_folio_lines_without_agreement() -> None:
    """— 3 —, · 4 ·, - 5 -: a line that is nothing but a folio in punctuation
    dies at the page edge even though no two pages agree on it."""
    pages = [
        "— 3 —\n正文甲的第一行。\n正文甲的第二行。",
        "· 4 ·\n正文乙的第一行。\n正文乙的第二行。",
        "- 5 -\n正文丙的第一行。\n正文丙的第二行。",
    ]
    assert knowledge_module._strip_page_furniture(pages) == [
        "正文甲的第一行。\n正文甲的第二行。",
        "正文乙的第一行。\n正文乙的第二行。",
        "正文丙的第一行。\n正文丙的第二行。",
    ]


def test_strip_page_furniture_keeps_a_single_page_whole() -> None:
    """One page has no repetition to observe, so nothing is judged."""
    page = "Jarvis 手册\n第 1 页\n正文。"
    assert knowledge_module._strip_page_furniture([page]) == [page]


@pytest.mark.asyncio
async def test_pdf_strips_a_header_that_carries_the_page_number(core: CoreServices) -> None:
    """The real-world leak: a Word export stamps 「产品手册 · 第 N 页」 on
    every page - verbatim never equal, folio-shaped and equal with digits
    folded. The body's own varying numbers (第 N 章) stay."""
    import pymupdf

    buffer = io.BytesIO()
    with pymupdf.open() as source:
        for index in range(1, 4):
            page = source.new_page()
            page.insert_text((72, 40), f"产品手册 · 第 {index} 页", fontsize=9, fontname="china-s")
            page.insert_text(
                (72, 120),
                f"登录模块第{index}章的接口约定：统一返回 JSON。",
                fontsize=11,
                fontname="china-s",
            )
        source.save(buffer)
    imported = await core.knowledge.import_items(
        [ImportItem(filename="手册.pdf", content=buffer.getvalue())], project_id=None
    )
    assert imported[-1]["status"] == "ready"
    content = imported[-1]["document"]["content"]

    assert "产品手册" not in content
    assert "登录模块第1章的接口约定" in content
    assert "登录模块第3章的接口约定" in content
